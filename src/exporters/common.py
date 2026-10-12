# -*- coding: utf-8 -*-
"""导出器共享的数据结构。

流程：
    DB (papers + paper_items + questions)
        ↓ paper_service.get_paper_detail / get_questions_by_ids
        ↓ build_paper_view()
    PaperView → 各格式 exporter → str / 文件
"""
from __future__ import annotations

from dataclasses import dataclass, field
from typing import Dict, List, Optional

import base64
import mimetypes
from pathlib import Path

from src.utils.file_utils import split_images


@dataclass
class ItemView:
    number: int             # 全卷连续编号，从 1 开始
    score: int
    question_id: int
    qtype: str
    stem: str
    options: List[dict]     # [{'key':'A','text':'...'}, ...]
    answer: str
    analysis: str
    image_list: List[str] = field(default_factory=list)           # 题干配图
    analysis_images: List[str] = field(default_factory=list)      # 解析配图
    missing: bool = False   # 题目已从题库移除，或从未取到


@dataclass
class SectionView:
    title: str
    items: List[ItemView] = field(default_factory=list)

    @property
    def total_score(self) -> int:
        return sum(it.score for it in self.items)

    @property
    def count(self) -> int:
        return len(self.items)


@dataclass
class PaperView:
    title: str
    subtitle: str
    exam_minutes: int
    total_score: int
    sections: List[SectionView]
    paper_id: Optional[int] = None


def build_paper_view(paper_detail: Dict, questions_map: Dict[int, object]) -> PaperView:
    """把 DB 结构转成导出视图。"""
    number = 0
    sections: List[SectionView] = []
    for si, sec in enumerate(paper_detail.get("sections", [])):
        sec_title = (sec.get("title") or "").strip() or f"大题{si + 1}"
        view = SectionView(title=sec_title)
        for it in sec.get("items", []):
            number += 1
            qid = int(it["question_id"])
            q = questions_map.get(qid)
            if q is None:
                view.items.append(ItemView(
                    number=number,
                    score=int(it.get("score", 0) or 0),
                    question_id=qid,
                    qtype="",
                    stem=f"（题目 #{qid} 已从题库中移除）",
                    options=[],
                    answer="",
                    analysis="",
                    image_list=[],
                    analysis_images=[],
                    missing=True,
                ))
                continue
            # 图片分两组：题干配图放题面，解析配图放答案与解析
            stem_images, analysis_images = split_images(q.image_list)
            view.items.append(ItemView(
                number=number,
                score=int(it.get("score", 0) or 0),
                question_id=qid,
                qtype=q.qtype or "",
                stem=q.stem or "",
                options=list(q.options_list or []),
                answer=q.answer or "",
                analysis=q.analysis or "",
                image_list=stem_images,
                analysis_images=analysis_images,
            ))
        sections.append(view)

    return PaperView(
        title=paper_detail.get("title") or "未命名试卷",
        subtitle=paper_detail.get("subtitle") or "",
        exam_minutes=int(paper_detail.get("exam_minutes") or 0),
        total_score=int(paper_detail.get("total_score") or 0),
        sections=sections,
        paper_id=paper_detail.get("id"),
    )

# ============================================================
# 正文分段（按 ![[img:N]] 标记把文字和图片拆开）
# ============================================================
def text_image_blocks(text: str, images) -> List[dict]:
    """把一段正文按插图标记拆成 [{kind:"text"/"image", ...}]。

    没有标记时就是「一段文字 + 末尾一串图片」，和以前完全一致；
    有标记时图片出现在标记所在的位置。

    Excel/HTML/Markdown 三种导出共用，保证导出的排版和界面里看到的一致。
    """
    from src.utils.rich_text import markdown_to_blocks

    return markdown_to_blocks(text or "", list(images or []))


# ============================================================
# 图片辅助
# ============================================================
def resolve_image_path(rel: str) -> Path | None:
    """相对 data/ 的路径 → 绝对 Path；文件不存在返回 None。

    入参可能是带归属前缀的路径（"analysis:images/xxx.png"），
    这里统一剥掉前缀再找文件。
    """
    if not rel:
        return None
    from src.config import DATA_DIR
    from src.utils.file_utils import real_image_path
    p = DATA_DIR / real_image_path(rel)
    return p if p.is_file() else None


def image_to_data_uri(rel: str) -> str | None:
    """读取图片并转成 data:image/...;base64,...，供 HTML / Markdown 内嵌。
    文件缺失或读取失败时返回 None。
    """
    p = resolve_image_path(rel)
    if p is None:
        return None
    mime, _ = mimetypes.guess_type(p.name)
    if not mime or not mime.startswith("image/"):
        mime = "image/png"
    try:
        raw = p.read_bytes()
    except OSError:
        return None
    b64 = base64.b64encode(raw).decode("ascii")
    return f"data:{mime};base64,{b64}"


# ============================================================
# Word 插图辅助
# ============================================================
def docx_image_source(path: Path):
    """返回可以直接交给 python-docx `add_picture()` 的图片来源。

    为什么需要它：python-docx 靠**固定的文件头签名**识别图片格式，它认的 JPEG 必须是
    `FF D8 FF E0 …JFIF` 或 `FF D8 FF E1 …Exif`（第 6 字节处是 JFIF/Exif）。
    下面两类图片会直接抛 UnrecognizedImageError，插不进 Word：
      1. webp —— python-docx 完全不支持，而上传时是允许 webp 的；
      2. 以 DQT 等其它标记开头的 JPEG（很多手机/截图/压缩工具产出的是 `FF D8 FF DB`）。
    这里先用原文件，python-docx 认不出来就用 Pillow 转成 PNG 再交给它。

    返回**字符串路径**（能用原文件）或 BytesIO（已转码）；都失败时返回 None。
    注意：python-docx 的 add_picture 只接受 str 或流，传 Path 会被当成流而报错，
    所以这里统一返回 str。
    """
    if not _docx_can_read(path):
        return transcode_to_png(path)
    return str(path)


def docx_add_picture(doc, path: Path, width_cm: float) -> bool:
    """把图片插入 Word 文档并居中，成功返回 True。

    自动兜住 python-docx 认不出的图片（webp、非标准文件头的 JPEG）。
    """
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm

    source = docx_image_source(path)
    if source is None:
        return False
    try:
        doc.add_picture(source, width=Cm(width_cm))
    except Exception:
        return False
    # add_picture 会新建一个段落，让它居中
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    return True


def _docx_can_read(path: Path) -> bool:
    """探测 python-docx 能否解析这张图片（不修改任何文档）。"""
    try:
        from docx.image.image import Image as _DocxImage
        _DocxImage.from_file(str(path))
        return True
    except Exception:
        return False


def transcode_to_png(path: Path):
    """用 Pillow 把图片转成 PNG 字节流（顺带按 EXIF 摆正方向）。失败返回 None。"""
    try:
        import io

        from PIL import Image, ImageOps
    except Exception:
        return None
    try:
        with Image.open(path) as im:
            im = ImageOps.exif_transpose(im) or im
            if im.mode not in ("RGB", "RGBA", "L"):
                im = im.convert("RGB")
            buf = io.BytesIO()
            im.save(buf, format="PNG")
        buf.seek(0)
        return buf
    except Exception:
        return None