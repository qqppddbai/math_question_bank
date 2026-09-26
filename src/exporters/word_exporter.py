# -*- coding: utf-8 -*-
"""Word 导出：python-docx，公式写成**原生 OMML**（Word 里可编辑的公式，不是 LaTeX 原文）。

管线：
    LaTeX --latex2mathml--> MathML --mathml2omml--> OMML --lxml--> <m:oMath>

上游两个库都有坑，这里全部兜住：

1. mathml2omml 返回的字符串**本身已经带 <m:oMath> 外壳**。再包一层就会
   变成 <m:oMath><m:oMath>… 嵌套结构（非法 OMML），Word 打开会提示
   "内容有问题，需要修复"。所以这里直接使用，不再包壳。
2. mathml2omml 0.0.2 把 <m:groupChrPr> 用 </m:groupChr> 闭合（复制粘贴笔误），
   产出非法 XML。\\vec \\bar \\overrightarrow \\widehat \\overbrace 等
   全部受影响，表现为公式整批变成错误提示文字。
3. latex2mathml 对 aligned / split / gather 等环境里的 & 对齐符会输出**裸 '&'**
   （非法 XML），而且 \\ 换行会退化成空 run（多行公式挤成一行）。
   这里把这类环境改写成 array：& 变列分隔、\\ 变行分隔，两边都能正确处理，
   在 Word 里就是真正的多行公式。

设计原则：任何一条公式转换失败都**只影响它自己** —— 退回该公式的 LaTeX
原文，绝不中断整份试卷的导出。
"""
from __future__ import annotations

import re
from functools import lru_cache
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from src.config import EXPORT_QUESTION_FONT_PT, QUESTION_FONT_NAME
from src.exporters.common import PaperView, resolve_image_path

# OMML 依赖是可选的：没装也不影响其它导出格式，只是公式退回 LaTeX 原文。
try:
    import latex2mathml.converter as _latex2mathml
    import mathml2omml as _mathml2omml

    OMML_READY = True
except ImportError:  # pragma: no cover - 取决于运行环境
    OMML_READY = False

# OMML 命名空间。mathml2omml 的输出不带 xmlns，parse_xml 前必须补上。
MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
# 公式 run 里要写 w:rPr（字体/字号），所以片段还得声明 w 命名空间。
WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

# 图片在 Word 里的最大显示宽度（厘米）。等比缩放。
WORD_IMAGE_WIDTH_CM = 9.0

# 题目显示统一字体：宋体 五号（10.5pt）。要改字号只改 src/config.py。
Q_FONT = QUESTION_FONT_NAME
Q_SIZE = EXPORT_QUESTION_FONT_PT

# 识别 $...$（行内）与 $$...$$（独立）公式。DOTALL 让公式可以跨行。
_MATH_RE = re.compile(r"(\$\$.*?\$\$|\$.*?\$)", re.DOTALL)

# ============================================================
# LaTeX 预处理
# ============================================================
# latex2mathml 把 aligned/align/split/gather 当普通分组处理：& 会变成裸 '&'，
# \\ 会变成 <mspace linebreak="newline"/>（mathml2omml 转出来是个空 run）。
# 换成 array 之后，两者都走 latex2mathml 本来就能正确处理的多行表格分支。
_ALIGN_ENVS = (
    "aligned", "alignedat", "split", "align", "align*",
    "gather", "gather*", "gathered", "multline", "multline*",
    "flalign", "flalign*", "eqnarray", "eqnarray*",
)
_ALIGN_ENV_RE = re.compile(
    r"\\begin\{(%s)\}(.*?)\\end\{\1\}" % "|".join(re.escape(e) for e in _ALIGN_ENVS),
    re.DOTALL,
)


def _align_env_to_array(match: re.Match) -> str:
    """把 aligned/align/split/... 环境改写成等价的 array 环境。"""
    body = match.group(2)
    rows = [r for r in re.split(r"\\\\", body) if r.strip()]
    # 列数 = 各行中 & 的最大个数 + 1
    ncol = max((r.count("&") + 1 for r in rows), default=1)
    # 第一列右对齐、其余左对齐，贴合 "a &= b" 的书写习惯
    spec = "r" + "l" * (ncol - 1)
    return "\\begin{array}{%s}%s\\end{array}" % (spec, body)


def _normalize_latex(latex: str) -> str:
    return _ALIGN_ENV_RE.sub(_align_env_to_array, latex)


# ============================================================
# MathML -> OMML
# ============================================================
# mathml2omml 0.0.2：<m:groupChrPr> 被 </m:groupChr> 错误闭合。
# 非贪婪匹配到第一个 </m:groupChr>，正是那个写错的闭合标签
# （groupChrPr 内部只会有 m:chr / m:pos / m:vertJc / m:ctrlPr）。
_GROUPCHR_BUG_RE = re.compile(r"<m:groupChrPr>(.*?)</m:groupChr>", re.DOTALL)

# latex2mathml 偶尔把 & 当普通字符输出成裸 '&'（XML 非法）。对齐符没有内容
# 语义，直接删掉；'&amp;' / '&#x3D;' 这类合法实体不受影响。
_BARE_AMP_RE = re.compile(r"&(?!(?:[A-Za-z][A-Za-z0-9]*|#\d+|#x[0-9A-Fa-f]+);)")

# ---- 修正"向量 / 上划线偏下" ----
# \vec \bar \overrightarrow \widetilde \overbrace 都转成 <m:groupChr>（把符号盖在基
# 字母上方）。这个对象还要一个 m:vertJc（Vertical Justification）说明基字母在框里
# 怎么摆，但 mathml2omml 从来不输出它。
# 微软规范 [MS-OE376] 7.1.2.42 明确说明："规范未规定 vertJc 的默认值，Word 按 top
# 处理"。于是基字母被顶到框顶，整个公式的基线被压到文字基线以下 —— 看起来就是
# "向量比旁边的字低一截"。
# texmath / pandoc 的做法（issue #98 里有 Word 存回来的对照 XML）是：符号在上
# (pos=top) 配 vertJc=bot，符号在下 (pos=bot) 配 vertJc=top。这里照做。
_GROUPCHR_PR_RE = re.compile(r"<m:groupChrPr>(.*?)</m:groupChrPr>", re.DOTALL)
_GROUPCHR_POS_RE = re.compile(r'<m:pos m:val="(top|bot)"/>')


def _add_groupchr_vertjc(omml: str) -> str:
    def _fix(match: re.Match) -> str:
        inner = match.group(1)
        if "<m:vertJc" in inner:            # 上游哪天自己补上了就不重复加
            return match.group(0)
        pos = _GROUPCHR_POS_RE.search(inner)
        if not pos:
            return match.group(0)
        # 符号在上 → 基字母贴底；符号在下 → 基字母贴顶
        vj = "bot" if pos.group(1) == "top" else "top"
        # CT_GroupChrPr 的子元素顺序是 chr, pos, vertJc, ctrlPr，插在 pos 后面
        inner = inner[:pos.end()] + '<m:vertJc m:val="%s"/>' % vj + inner[pos.end():]
        return "<m:groupChrPr>%s</m:groupChrPr>" % inner

    return _GROUPCHR_PR_RE.sub(_fix, omml)


# ---- 公式字体 / 字号 ----
# 数学 run 不给 w:rPr 的话会继承 docDefaults（python-docx 模板是 11pt 主题字体），
# 比正文的宋体五号（10.5pt）大一号，公式看起来"发胖"。Word 自己导出公式时会写
# <m:r><w:rPr><w:rFonts w:ascii="Cambria Math" .../></w:rPr><m:t>…
# 这里照做，并让字号跟调用方传来的正文字号一致。
MATH_FONT = "Cambria Math"
_M_R_RE = re.compile(r"<m:r>(<m:rPr>.*?</m:rPr>)?", re.DOTALL)


def _math_rpr(size_pt: float) -> str:
    half = int(round(size_pt * 2))          # OOXML 字号单位是半磅
    return (
        '<w:rPr><w:rFonts w:ascii="{f}" w:hAnsi="{f}" w:cs="{f}"/>'
        '<w:sz w:val="{s}"/><w:szCs w:val="{s}"/></w:rPr>'
    ).format(f=MATH_FONT, s=half)


@lru_cache(maxsize=4096)
def _omml_string(latex: str, size_pt: float = Q_SIZE) -> str:
    """LaTeX → '<m:oMath>…</m:oMath>' 字符串（已修上游 bug、已补命名空间）。

    公式高度重复（一道题里同一个式子可能出现多次），这里缓存转换结果。
    """
    mathml = _latex2mathml.convert(_normalize_latex(latex))
    mathml = _BARE_AMP_RE.sub("", mathml)
    omml = _mathml2omml.convert(mathml)
    omml = _GROUPCHR_BUG_RE.sub(r"<m:groupChrPr>\1</m:groupChrPr>", omml)
    omml = _add_groupchr_vertjc(omml)
    # 给每个 m:r 补 w:rPr；有 m:rPr 时必须排在它后面（CT_R 的元素顺序）
    rpr = _math_rpr(size_pt)
    omml = _M_R_RE.sub(lambda m: "<m:r>" + (m.group(1) or "") + rpr, omml)
    if "xmlns:m=" not in omml:
        omml = omml.replace(
            "<m:oMath",
            '<m:oMath xmlns:m="%s" xmlns:w="%s"' % (MATH_NS, WORD_NS),
            1,
        )
    return omml


def _omath_element(latex: str, size_pt: float = Q_SIZE):
    """行内公式：<m:oMath> 元素，直接插进 w:p，与文本 run 交错。"""
    return parse_xml(_omml_string(latex, size_pt))


def _omathpara_element(latex: str, size_pt: float = Q_SIZE):
    """独立公式：<m:oMathPara>（居中）。

    注意 m:oMathPara 是块级元素，必须独占一个 w:p，所以调用方会为它单开段落。
    """
    return parse_xml(
        '<m:oMathPara xmlns:m="%s" xmlns:w="%s">'
        '<m:oMathParaPr><m:jc m:val="center"/></m:oMathParaPr>'
        "%s"
        "</m:oMathPara>" % (MATH_NS, WORD_NS, _omml_string(latex, size_pt))
    )


# ============================================================
# 文本 / 公式混排
# ============================================================
def _style_run(run, *, name: str = Q_FONT, size: float = Q_SIZE, bold: bool = False,
               italic: bool = False):
    run.bold = bold
    run.italic = italic
    run.font.name = name
    run.font.size = Pt(size)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), name)


def _add_plain_text(paragraph, text: str, *, name=Q_FONT, size=Q_SIZE,
                    bold=False, italic=False):
    """添加普通文本，把 \n 转成 Word 的换行（run.add_break）。"""
    if not text:
        return
    for i, part in enumerate(text.replace("\r\n", "\n").split("\n")):
        if i > 0:
            br = paragraph.add_run()
            _style_run(br, name=name, size=size, bold=bold, italic=italic)
            br.add_break()
        run = paragraph.add_run(part)
        _style_run(run, name=name, size=size, bold=bold, italic=italic)


def _is_empty_paragraph(paragraph) -> bool:
    """段落除 w:pPr 外没有任何内容。"""
    return all(child.tag == qn("w:pPr") for child in paragraph._element)


def emit_text(doc, text: str, *, name=Q_FONT, size=Q_SIZE, bold=False,
               italic=False, paragraph=None, indent_cm=None):
    """把「普通文本 + $行内公式$ + $$独立公式$$」混排写进文档。

    - 行内公式：<m:oMath> 与文本 run 按原文顺序交错在同一段落里；
    - 独立公式：单开一个居中段落，只放 <m:oMathPara>（OMML 规范要求独占段落），
      公式前后的换行会被吃掉，避免多出空行；
    - paragraph 传 None 时自动新建段落；indent_cm 用于后续新建段落的左缩进。

    返回最后一个段落（可能为 None：文本以独立公式结尾时不再补空段落）。
    """
    if not text:
        return paragraph

    def new_para():
        p = doc.add_paragraph()
        if indent_cm is not None:
            p.paragraph_format.left_indent = Cm(indent_cm)
        return p

    # 依赖缺失：整段按原文输出，内容不丢
    if not OMML_READY:
        target = paragraph if paragraph is not None else new_para()
        _add_plain_text(target, text, name=name, size=size, bold=bold, italic=italic)
        return target

    parts = _MATH_RE.split(text)
    # re.split 带捕获组：偶数下标是普通文本，奇数下标是公式片段
    for i in range(1, len(parts), 2):
        if not (parts[i].startswith("$$") and parts[i].endswith("$$")):
            continue
        # 独立公式自成一段，紧邻的换行是多余的，去掉以免多出空行
        if i - 1 >= 0:
            parts[i - 1] = parts[i - 1].rstrip("\r\n")
        if i + 1 < len(parts):
            parts[i + 1] = parts[i + 1].lstrip("\r\n")

    cur = paragraph
    for i, chunk in enumerate(parts):
        if not chunk:
            continue

        # ---- 普通文本 ----
        if i % 2 == 0:
            if cur is None:
                cur = new_para()
            _add_plain_text(cur, chunk, name=name, size=size, bold=bold, italic=italic)
            continue

        # ---- 公式 ----
        is_display = chunk.startswith("$$") and chunk.endswith("$$")
        latex = (chunk[2:-2] if is_display else chunk[1:-1]).strip()
        if not latex:
            continue
        try:
            if is_display:
                target = cur if (cur is not None and _is_empty_paragraph(cur)) else new_para()
                target.alignment = WD_ALIGN_PARAGRAPH.CENTER
                target._element.append(_omathpara_element(latex, size))
                cur = None      # 公式后面的文字另起一段，没有就不补空段落
            else:
                if cur is None:
                    cur = new_para()
                cur._element.append(_omath_element(latex, size))
        except Exception:
            # 单条公式转换失败：退回 LaTeX 原文，保证内容不丢、导出不中断
            if cur is None:
                cur = new_para()
            _add_plain_text(cur, chunk, name=name, size=size, bold=bold, italic=italic)

    return cur


# ============================================================
# 版面辅助
# ============================================================
def _section_meta(sec) -> str:
    if not sec.items:
        return ""
    scores = {it.score for it in sec.items}
    if len(scores) == 1:
        s = next(iter(scores))
        return f"（共 {sec.count} 题，每题 {s} 分，共 {sec.total_score} 分）"
    return f"（共 {sec.count} 题，共 {sec.total_score} 分）"


def _add_images(doc: Document, image_list):
    """把题目图片按顺序插入，居中，等比缩放到 WORD_IMAGE_WIDTH_CM 宽。"""
    for rel in image_list or []:
        p = resolve_image_path(rel)
        if p is None:
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Cm(0.8)
            run = para.add_run("［图片丢失］")
            _style_run(run, name="宋体", size=10, italic=True)
            continue
        try:
            doc.add_picture(str(p), width=Cm(WORD_IMAGE_WIDTH_CM))
            # add_picture 会新建一个段落，让它居中
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        except Exception:
            # 单个图片失败不影响整份试卷
            para = doc.add_paragraph()
            run = para.add_run("［图片插入失败］")
            _style_run(run, name="宋体", size=10, italic=True)


def _add_labeled_block(doc: Document, label: str, text: str):
    """「答案：/解析：」+ 混排正文（缩进 0.8cm，公式按行内处理）。"""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.8)
    run = p.add_run(label)
    _style_run(run, name=Q_FONT, size=Q_SIZE, bold=True)
    emit_text(doc, text, name=Q_FONT, size=Q_SIZE, paragraph=p, indent_cm=0.8)


# ============================================================
# 入口
# ============================================================
def export_word(view: PaperView, filepath: Path, *, include_answer: bool = True,
                include_analysis: bool = True, answers_at_end: bool = False,
                include_images: bool = True) -> None:
    doc = Document()

    for sec in doc.sections:
        sec.top_margin = Cm(2.0)
        sec.bottom_margin = Cm(2.0)
        sec.left_margin = Cm(2.5)
        sec.right_margin = Cm(2.5)

    # 标题
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    emit_text(doc, view.title, name="黑体", size=16, bold=True, paragraph=p)

    if view.subtitle:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        emit_text(doc, view.subtitle, name="楷体", size=12, paragraph=p)

    meta = []
    if view.exam_minutes:
        meta.append(f"考试时间：{view.exam_minutes} 分钟")
    meta.append(f"满分：{view.total_score} 分")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    emit_text(doc, "　".join(meta), name="宋体", size=11, paragraph=p)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    emit_text(doc, "姓名：__________　学校：__________　得分：__________",
               name="宋体", size=11, paragraph=p)

    # 正文
    for sec in view.sections:
        p = doc.add_paragraph()
        emit_text(doc, sec.title, name="黑体", size=13, bold=True, paragraph=p)
        m = _section_meta(sec)
        if m:
            p = doc.add_paragraph()
            emit_text(doc, m, name="楷体", size=10, italic=True, paragraph=p)

        for it in sec.items:
            p = doc.add_paragraph()
            run = p.add_run(f"{it.number}.")
            _style_run(run, name=Q_FONT, size=Q_SIZE, bold=True)
            run = p.add_run(f"（{it.score} 分） ")
            _style_run(run, name=Q_FONT, size=Q_SIZE)
            emit_text(doc, it.stem, name=Q_FONT, size=Q_SIZE, paragraph=p)

            if include_images and it.image_list:
                _add_images(doc, it.image_list)

            for o in it.options:
                text = (o.get("text") or "").strip()
                if not text:
                    continue
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.8)
                run = p.add_run(f"{o.get('key','')}. ")
                _style_run(run, name=Q_FONT, size=Q_SIZE, bold=True)
                emit_text(doc, text, name=Q_FONT, size=Q_SIZE, paragraph=p,
                           indent_cm=0.8)

            if not answers_at_end and (include_answer or include_analysis):
                if include_answer and it.answer:
                    _add_labeled_block(doc, "答案：", it.answer)
                if include_analysis and it.analysis:
                    _add_labeled_block(doc, "解析：", it.analysis)

    # 卷末答案
    if answers_at_end and (include_answer or include_analysis):
        doc.add_page_break()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        emit_text(doc, "参考答案与解析", name="黑体", size=16, bold=True, paragraph=p)

        for sec in view.sections:
            p = doc.add_paragraph()
            emit_text(doc, sec.title, name="黑体", size=13, bold=True, paragraph=p)
            for it in sec.items:
                p = doc.add_paragraph()
                run = p.add_run(f"{it.number}.")
                _style_run(run, name=Q_FONT, size=Q_SIZE, bold=True)

                if include_images and it.image_list:
                    _add_images(doc, it.image_list)

                if include_answer and it.answer:
                    _add_labeled_block(doc, "答案：", it.answer)
                if include_analysis and it.analysis:
                    _add_labeled_block(doc, "解析：", it.analysis)

    doc.save(str(filepath))
