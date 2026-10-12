# -*- coding: utf-8 -*-
"""学生档案导出：**错题记录** 与 **学习图片** 分开导出（Word / HTML / Markdown）。

    DB 数据 → build_student_view() → StudentView → 各格式 exporter

错题记录的版式约定：
- 正文部分只放题干 / 题目图片 / 学生错答 / 错因分析（含章节、知识点）；
- **答案与解析统一放到文末，并另起一页**，方便先让学生重做再看答案；
- 每题标注章节与知识点。
"""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Dict, List

from src.config import (EXPORT_QUESTION_FONT_PT, QUESTION_FONT_NAME,
                        QUESTION_FONT_STACK)
from src.exporters.common import (docx_add_picture, image_to_data_uri,
                                  resolve_image_path)
from src.exporters.word_exporter import emit_text

# 与 student_page 保持一致；放在 exporter 里避免循环导入
MASTERY_LABELS = {
    0: "❌ 未掌握",
    1: "⚠️ 部分掌握",
    2: "✅ 基本掌握",
    3: "🌟 已掌握",
}

WORD_IMAGE_WIDTH_CM = 9.0

# 题目显示统一字体：宋体 五号（10.5pt）。要改字号只改 src/config.py。
Q_FONT = QUESTION_FONT_NAME
Q_SIZE = EXPORT_QUESTION_FONT_PT

ANSWER_SECTION_TITLE = "参考答案与解析"


# ============================================================
# 视图结构
# ============================================================
@dataclass
class MistakeView:
    number: int
    question_id: int
    stem: str
    answer: str
    analysis: str
    wrong_answer: str
    error_reason: str
    mastery: int
    created_at: datetime | None
    chapter: str = ""
    knowledge: str = ""
    qtype: str = ""
    images: List[str] = field(default_factory=list)      # 相对 data/ 的路径


@dataclass
class LooseImageView:
    path: str
    caption: str = ""


@dataclass
class StudentView:
    name: str
    class_name: str
    school: str
    notes: str
    mistakes: List[MistakeView] = field(default_factory=list)
    loose_images: List[LooseImageView] = field(default_factory=list)

    @property
    def mistakes_title(self) -> str:
        return f"{self.name or '学生'} · 错题记录"

    @property
    def pictures_title(self) -> str:
        return f"{self.name or '学生'} · 学习图片"

    @property
    def info_line(self) -> str:
        """学校 / 班级一行（没有年级）。"""
        parts = []
        if self.school:
            parts.append(f"学校：{self.school}")
        if self.class_name:
            parts.append(f"班级：{self.class_name}")
        return "　".join(parts)


def build_student_view(student, mistakes: List[Dict],
                       loose_files: List[Dict]) -> StudentView:
    """从 service 层拿到的 dict 结构构造导出视图。"""
    mv: List[MistakeView] = []
    for i, m in enumerate(mistakes, start=1):
        mv.append(MistakeView(
            number=i,
            question_id=m.get("question_id") or 0,
            stem=m.get("snapshot_stem") or "",
            answer=m.get("snapshot_answer") or "",
            analysis=m.get("snapshot_analysis") or "",
            wrong_answer=m.get("wrong_answer") or "",
            error_reason=m.get("error_reason") or "",
            mastery=int(m.get("mastery", 0) or 0),
            created_at=m.get("created_at"),
            chapter=m.get("chapter") or "",
            knowledge=m.get("knowledge") or "",
            qtype=m.get("qtype") or "",
            images=[f["path"] for f in m.get("files", []) if f.get("path")],
        ))

    loose = [LooseImageView(path=f["path"], caption=f.get("caption") or "")
             for f in loose_files if f.get("path")]

    return StudentView(
        name=student.name or "",
        school=getattr(student, "school", "") or "",
        class_name=student.class_name or "",
        notes=student.notes or "",
        mistakes=mv,
        loose_images=loose,
    )


def _mistake_meta_line(m: MistakeView) -> str:
    """「章节：…　知识点：…」一行（没有就返回空串）。"""
    parts = []
    if m.chapter:
        parts.append(f"章节：{m.chapter}")
    if m.knowledge:
        parts.append(f"知识点：{m.knowledge}")
    return "　".join(parts)


# ============================================================
# 公共小工具
# ============================================================
def _ml(text: str) -> str:
    return (text or "").replace("\r\n", "\n").replace("\n", "  \n")


def _emit_images_md(lines, image_list, missing_hint="［图片丢失］"):
    for rel in image_list or []:
        uri = image_to_data_uri(rel)
        if uri:
            lines.append(f"![]({uri})")
            lines.append("")
        else:
            lines.append(f"_{missing_hint}_")
            lines.append("")


def _esc(text: str) -> str:
    import html as _htmlmod
    return _htmlmod.escape(text or "", quote=False)


def _render_images_html(image_list) -> str:
    if not image_list:
        return ""
    out = ['<div class="images">']
    for rel in image_list:
        uri = image_to_data_uri(rel)
        if uri:
            out.append(f'<img src="{uri}" alt="图片">')
        else:
            out.append('<div class="missing-image">［图片丢失］</div>')
    out.append('</div>')
    return "".join(out)


# ============================================================
# 错题记录 —— Markdown
# ============================================================
def export_mistakes_markdown(view: StudentView) -> str:
    lines = [f"# {view.mistakes_title}", ""]
    if view.info_line:
        lines += [view.info_line, ""]
    if view.notes:
        lines += [f"*备注：{view.notes}*", ""]
    lines += [f"共 {len(view.mistakes)} 道错题", "", "---", ""]

    if not view.mistakes:
        lines += ["_暂无错题。_", ""]

    for m in view.mistakes:
        lines.append(f"### {m.number}. 题目 #{m.question_id}")
        created = m.created_at.strftime("%Y-%m-%d") if m.created_at else "—"
        lines.append(f"*掌握程度：{MASTERY_LABELS.get(m.mastery, '—')}　记录于 {created}*")
        lines.append("")
        meta_line = _mistake_meta_line(m)
        if meta_line:
            lines += [meta_line, ""]

        lines += ["**题干**", "", _ml(m.stem) or "_（无）_", ""]

        if m.images:
            _emit_images_md(lines, m.images)

        if m.wrong_answer:
            lines += ["**学生错答**", "", "> " + _ml(m.wrong_answer), ""]
        if m.error_reason:
            lines += ["**错因分析**", "", _ml(m.error_reason), ""]
        lines += ["---", ""]

    # ---- 答案与解析：文末另起一页 ----
    lines += ['<div style="page-break-before: always"></div>', ""]
    lines += [f"# {ANSWER_SECTION_TITLE}", ""]
    for m in view.mistakes:
        lines.append(f"**{m.number}. 题目 #{m.question_id}**")
        lines.append("")
        if m.answer:
            lines += ["答案：", "", _ml(m.answer), ""]
        else:
            lines += ["答案：_（无）_", ""]
        if m.analysis:
            lines += ["解析：", "", _ml(m.analysis), ""]
    return "\n".join(lines)


# ============================================================
# 错题记录 —— HTML
# ============================================================
def export_mistakes_html(view: StudentView) -> str:
    body = [f'<h1 class="title">{_esc(view.mistakes_title)}</h1>']
    if view.info_line:
        body.append(f'<div class="meta">{_esc(view.info_line)}</div>')
    if view.notes:
        body.append(f'<div class="notes">备注：{_esc(view.notes)}</div>')
    body.append(f'<div class="meta">共 {len(view.mistakes)} 道错题</div>')

    body.append('<h2>错题记录</h2>')
    if not view.mistakes:
        body.append('<p><i>暂无错题。</i></p>')

    for m in view.mistakes:
        created = m.created_at.strftime("%Y-%m-%d") if m.created_at else "—"
        body.append('<div class="mistake">')
        body.append(
            f'<div class="m-head"><b>{m.number}. 题目 #{m.question_id}</b>'
            f'<span class="m-meta">掌握程度：{_esc(MASTERY_LABELS.get(m.mastery, "—"))}'
            f'　记录于 {created}</span></div>'
        )
        meta_line = _mistake_meta_line(m)
        if meta_line:
            body.append(f'<div class="m-chapter">{_esc(meta_line)}</div>')
        body.append('<div class="block"><div class="label">题干</div>'
                    f'<div class="stem">{_esc(m.stem) or "（无）"}</div></div>')
        if m.images:
            body.append(_render_images_html(m.images))
        if m.wrong_answer:
            body.append('<div class="block wrong"><div class="label">学生错答</div>'
                        f'<div class="text">{_esc(m.wrong_answer)}</div></div>')
        if m.error_reason:
            body.append('<div class="block reason"><div class="label">错因分析</div>'
                        f'<div class="text">{_esc(m.error_reason)}</div></div>')
        body.append('</div>')

    # ---- 答案与解析：文末另起一页 ----
    body.append('<div class="page-break"></div>')
    body.append(f'<h1 class="title">{ANSWER_SECTION_TITLE}</h1>')
    for m in view.mistakes:
        body.append('<div class="mistake">')
        body.append(f'<div class="m-head"><b>{m.number}. 题目 #{m.question_id}</b></div>')
        body.append('<div class="block"><div class="label">答案</div>'
                    f'<div class="answer">{_esc(m.answer) or "（无）"}</div></div>')
        if m.analysis:
            body.append('<div class="block"><div class="label">解析</div>'
                        f'<div class="text">{_esc(m.analysis)}</div></div>')
        body.append('</div>')

    return _html_document(
        title=view.mistakes_title,
        extra_css=".m-chapter { color:#666; font-size: 12px; margin-bottom: 0.5em; }",
        content="\n".join(body),
    )


# ============================================================
# 学习图片 —— Markdown / HTML
# ============================================================
def export_pictures_markdown(view: StudentView) -> str:
    lines = [f"# {view.pictures_title}", ""]
    if view.info_line:
        lines += [view.info_line, ""]
    if view.notes:
        lines += [f"*备注：{view.notes}*", ""]
    lines += [f"共 {len(view.loose_images)} 张图片", "", "---", ""]

    if not view.loose_images:
        lines += ["_暂无学习图片。_", ""]

    for i, img in enumerate(view.loose_images, start=1):
        if img.caption:
            lines += [f"**{i}. {img.caption}**", ""]
        else:
            lines += [f"**{i}.**", ""]
        _emit_images_md(lines, [img.path])
    return "\n".join(lines)


def export_pictures_html(view: StudentView) -> str:
    body = [f'<h1 class="title">{_esc(view.pictures_title)}</h1>']
    if view.info_line:
        body.append(f'<div class="meta">{_esc(view.info_line)}</div>')
    if view.notes:
        body.append(f'<div class="notes">备注：{_esc(view.notes)}</div>')
    body.append(f'<div class="meta">共 {len(view.loose_images)} 张图片</div>')

    if not view.loose_images:
        body.append('<p><i>暂无学习图片。</i></p>')

    for i, img in enumerate(view.loose_images, start=1):
        body.append('<div class="picture">')
        body.append(f'<div class="caption"><b>{i}.</b> {_esc(img.caption)}</div>')
        body.append(_render_images_html([img.path]))
        body.append('</div>')

    return _html_document(
        title=view.pictures_title,
        extra_css=".picture { margin: 1.4em 0; page-break-inside: avoid; }",
        content="\n".join(body),
    )


# ============================================================
# HTML 外壳
# ============================================================
KATEX_VERSION = "0.16.9"
KATEX_CSS = f"https://cdn.jsdelivr.net/npm/katex@{KATEX_VERSION}/dist/katex.min.css"
KATEX_JS = f"https://cdn.jsdelivr.net/npm/katex@{KATEX_VERSION}/dist/katex.min.js"
KATEX_AUTO = f"https://cdn.jsdelivr.net/npm/katex@{KATEX_VERSION}/dist/contrib/auto-render.min.js"


def _html_document(*, title: str, content: str, extra_css: str = "") -> str:
    return f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="UTF-8">
<title>{_esc(title)}</title>
<link rel="stylesheet" href="{KATEX_CSS}">
<style>
  body {{
    font-family: {QUESTION_FONT_STACK};
    max-width: 820px;
    margin: 2em auto;
    line-height: 1.85;
    font-size: {Q_SIZE}pt;   /* 题目正文：宋体 五号 */
    color: #000;
    padding: 0 1em;
  }}
  h1.title {{ text-align:center; font-size: 22px; margin-bottom: 0.4em; }}
  .meta {{ text-align:center; color:#333; margin-bottom: 0.3em; }}
  .notes {{ text-align:center; color:#666; font-size: 13px; margin-bottom: 0.6em; }}
  h2 {{
    font-size: 17px;
    border-bottom: 1px solid #bbb;
    padding-bottom: 4px;
    margin-top: 2em;
    margin-bottom: 0.6em;
  }}
  .mistake {{
    margin: 1.6em 0;
    padding: 0.8em 1em;
    border: 1px solid #e0e0e0;
    border-radius: 6px;
    page-break-inside: avoid;
  }}
  .m-head {{ margin-bottom: 0.6em; }}
  .m-meta {{ color:#666; font-size: 12px; margin-left: 1em; }}
  .block {{ margin: 0.5em 0; }}
  .block .label {{ font-weight: bold; color: #444; font-size: inherit; }}
  .stem {{ white-space: pre-wrap; word-wrap: break-word; }}
  .answer {{ white-space: pre-wrap; word-wrap: break-word; color: #333; }}
  /* 学生错答 / 错因分析：同样保留原始换行 */
  .text {{ white-space: pre-wrap; word-wrap: break-word; }}
  .block.wrong {{
    padding: 0.5em 0.8em;
    background: #fff2f2;
    border-left: 3px solid #e55;
  }}
  .block.reason {{
    padding: 0.5em 0.8em;
    background: #f6f6ff;
    border-left: 3px solid #77a;
  }}
  .images {{ margin: 0.6em 0 0.6em 1em; }}
  .images img {{ max-width: 100%; height: auto; display: block; margin: 0.4em 0; }}
  .missing-image {{
    color: #999; font-size: 13px; font-style: italic;
    padding: 4px 8px; border: 1px dashed #ccc; display: inline-block;
  }}
  .caption {{ margin-top: 1em; color: #333; }}
  /* 答案与解析另起一页 */
  .page-break {{ page-break-before: always; break-before: page; }}
  {extra_css}
  @media print {{
    body {{ max-width: 100%; margin: 0; padding: 0 15mm; font-size: {Q_SIZE}pt; }}
    .images img {{ max-width: 75%; }}
    .mistake {{ border-color: #ccc; }}
  }}
</style>
</head>
<body>
{content}
<script src="{KATEX_JS}"></script>
<script src="{KATEX_AUTO}"></script>
<script>
  renderMathInElement(document.body, {{
    delimiters: [
      {{left: '$$', right: '$$', display: true}},
      {{left: '$', right: '$', display: false}},
      {{left: '\\\\(', right: '\\\\)', display: false}},
      {{left: '\\\\[', right: '\\\\]', display: true}}
    ],
    throwOnError: false
  }});
</script>
</body>
</html>
"""


# ============================================================
# Word 公共小工具
# ============================================================
def _docx_setup():
    from docx import Document
    from docx.shared import Cm

    doc = Document()
    for sec in doc.sections:
        sec.top_margin = Cm(2.0)
        sec.bottom_margin = Cm(2.0)
        sec.left_margin = Cm(2.5)
        sec.right_margin = Cm(2.5)
    return doc


def _docx_style_run(run, *, name=Q_FONT, size=Q_SIZE, bold=False, italic=False):
    from docx.oxml.ns import qn
    from docx.shared import Pt

    run.bold = bold
    run.italic = italic
    run.font.name = name
    run.font.size = Pt(size)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), name)


def _docx_add_text(doc, paragraph, text, *, name=Q_FONT, size=Q_SIZE,
                   bold=False, italic=False):
    """复用试卷导出的混排逻辑：$...$ 写成真正的 OMML 公式。"""
    li = paragraph.paragraph_format.left_indent
    emit_text(doc, text, name=name, size=size, bold=bold, italic=italic,
              paragraph=paragraph, indent_cm=li.cm if li is not None else None)


def _docx_add_images(doc, image_list, indent_cm: float = 1.0):
    """插入图片。webp / 非标准文件头的 JPEG 会自动转成 PNG（见 docx_add_picture）。"""
    from docx.shared import Cm

    for rel in image_list or []:
        p = resolve_image_path(rel)
        if p is not None and docx_add_picture(doc, p, WORD_IMAGE_WIDTH_CM):
            continue

        para = doc.add_paragraph()
        para.paragraph_format.left_indent = Cm(indent_cm)
        label = "［图片插入失败］" if p is not None else "［图片丢失］"
        _docx_style_run(para.add_run(label), size=10, italic=True)


def _docx_head(doc, title: str, view: StudentView, subtitle_count: str):
    """标题 + 学校/班级 + 备注 + 数量说明。"""
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _docx_add_text(doc, p, title, name="黑体", size=16, bold=True)

    if view.info_line:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _docx_add_text(doc, p, view.info_line, name="宋体", size=11)

    if view.notes:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _docx_add_text(doc, p, f"备注：{view.notes}", name="楷体", size=10, italic=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _docx_add_text(doc, p, subtitle_count, name="宋体", size=11)


# ============================================================
# 错题记录 —— Word
# ============================================================
def export_mistakes_word(view: StudentView, filepath: Path) -> None:
    from docx.shared import Cm

    doc = _docx_setup()
    _docx_head(doc, view.mistakes_title, view, f"共 {len(view.mistakes)} 道错题")

    p = doc.add_paragraph()
    _docx_add_text(doc, p, "错题记录", name="黑体", size=13, bold=True)

    if not view.mistakes:
        p = doc.add_paragraph()
        _docx_add_text(doc, p, "暂无错题。", name="宋体", size=11, italic=True)

    for m in view.mistakes:
        # 题号
        p = doc.add_paragraph()
        _docx_style_run(p.add_run(f"{m.number}. 题目 #{m.question_id}"), bold=True)

        # 章节 / 知识点
        meta_line = _mistake_meta_line(m)
        if meta_line:
            p = doc.add_paragraph()
            _docx_add_text(doc, p, meta_line, name="楷体", size=10, italic=True)

        # 掌握程度 / 记录时间
        created = m.created_at.strftime("%Y-%m-%d") if m.created_at else "—"
        p = doc.add_paragraph()
        _docx_add_text(doc, p, f"掌握程度：{MASTERY_LABELS.get(m.mastery, '—')}　"
                               f"记录于 {created}",
                       name="楷体", size=10, italic=True)

        # 题干
        p = doc.add_paragraph()
        _docx_style_run(p.add_run("题干："), bold=True)
        _docx_add_text(doc, p, m.stem or "（无）")

        if m.images:
            _docx_add_images(doc, m.images)

        if m.wrong_answer:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            _docx_style_run(p.add_run("学生错答："), bold=True)
            _docx_add_text(doc, p, m.wrong_answer)

        if m.error_reason:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            _docx_style_run(p.add_run("错因分析："), bold=True)
            _docx_add_text(doc, p, m.error_reason)

    # ---- 答案与解析：另起一页 ----
    doc.add_page_break()
    p = doc.add_paragraph()
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _docx_add_text(doc, p, ANSWER_SECTION_TITLE, name="黑体", size=16, bold=True)

    for m in view.mistakes:
        p = doc.add_paragraph()
        _docx_style_run(p.add_run(f"{m.number}. 题目 #{m.question_id}"), bold=True)

        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        _docx_style_run(p.add_run("答案："), bold=True)
        _docx_add_text(doc, p, m.answer or "（无）")

        if m.analysis:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            _docx_style_run(p.add_run("解析："), bold=True)
            _docx_add_text(doc, p, m.analysis)

    doc.save(str(filepath))


# ============================================================
# 学习图片 —— Word
# ============================================================
def export_pictures_word(view: StudentView, filepath: Path) -> None:
    doc = _docx_setup()
    _docx_head(doc, view.pictures_title, view,
               f"共 {len(view.loose_images)} 张图片")

    if not view.loose_images:
        p = doc.add_paragraph()
        _docx_add_text(doc, p, "暂无学习图片。", name="宋体", size=11, italic=True)

    for i, img in enumerate(view.loose_images, start=1):
        p = doc.add_paragraph()
        _docx_style_run(p.add_run(f"{i}."), bold=True)
        if img.caption:
            _docx_add_text(doc, p, img.caption)
        _docx_add_images(doc, [img.path], indent_cm=0.0)

    doc.save(str(filepath))
