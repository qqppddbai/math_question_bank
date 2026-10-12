# -*- coding: utf-8 -*-
"""自测 C：导出（HTML / Markdown / Word）时题干图与解析图各归各位。"""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[3]   # 项目根目录（tools/verify/checks → 项目根）
sys.path.insert(0, str(ROOT))

from src.exporters.common import build_paper_view  # noqa: E402
from src.exporters.html_exporter import export_html  # noqa: E402
from src.exporters.markdown_exporter import export_markdown  # noqa: E402
from src.exporters.word_exporter import export_word  # noqa: E402
from src.services import paper_service as ps  # noqa: E402
from src.services import question_service as qs  # noqa: E402

results = []


def check(name, ok, detail=""):
    results.append((name, bool(ok), detail))
    print(f"[{'PASS' if ok else 'FAIL'}] {name}" + (f"  -> {detail}" if detail else ""))


# 用 #3（题干图 + 解析图）组一份试卷
qs.update_question(3, {
    "image_list": ["images/both_stem.png", "analysis:images/both_analysis.png"],
})
paper_id = ps.create_paper(
    title="导出验证卷", subtitle="", exam_minutes=0,
    sections=[{"title": "一、解答题",
               "items": [{"question_id": 3, "score": 0}]}],
)
detail = ps.get_paper_detail(paper_id)
qmap = ps.get_questions_by_ids([3])
view = build_paper_view(detail, qmap)

item = view.sections[0].items[0]
check("ItemView 已把两组图分开",
      item.image_list == ["images/both_stem.png"]
      and item.analysis_images == ["images/both_analysis.png"],
      f"stem={item.image_list} analysis={item.analysis_images}")

# ---------------- HTML ----------------
html = export_html(view, include_images=True, answers_at_end=False)
stem_pos = html.find('class="stem"')
ans_pos = html.find('<div class="answer-block">')   # 注意别匹配到 CSS 里的 .answer-block
first_img = html.find("data:image")
last_img = html.rfind("data:image")
check("HTML：两张图都内嵌", html.count("data:image") == 2,
      f"图片数={html.count('data:image')}")
check("HTML：题干图在题面（答案块之前）",
      stem_pos < first_img < ans_pos, f"stem={stem_pos} img={first_img} ans={ans_pos}")
check("HTML：解析图在答案块之后", last_img > ans_pos,
      f"last_img={last_img} ans={ans_pos}")

# ---------------- HTML：带 ![[img:N]] 标记时要插到文字之间 ----------------
qs.update_question(3, {
    "stem": "题干前半段。\n\n![[img:1]]\n\n题干后半段。",
    "analysis": "解析前半段。\n\n![[img:1]]\n\n解析后半段。",
    "image_list": ["images/both_stem.png", "analysis:images/both_analysis.png"],
})
view2 = build_paper_view(ps.get_paper_detail(paper_id), ps.get_questions_by_ids([3]))
html2 = export_html(view2, include_images=True, answers_at_end=False)
p_head = html2.find("题干前半段。")
p_img = html2.find("data:image")
p_tail = html2.find("题干后半段。")
check("HTML：题干图插在两段文字之间（不再堆到末尾）",
      -1 < p_head < p_img < p_tail, f"{p_head} {p_img} {p_tail}")
check("HTML：标记本身不会残留到导出结果里",
      "![[img:1]]" not in html2, "还能看到标记原文" if "![[img:1]]" in html2 else "")

a_head = html2.find("解析前半段。")
a_img = html2.rfind("data:image")
check("HTML：解析图也在解析文字里", -1 < a_head < a_img)

# ---------------- Markdown：同上 ----------------
md2 = export_markdown(view2, include_images=True, answers_at_end=False)
check("Markdown：题干图插在两段文字之间",
      -1 < md2.find("题干前半段。") < md2.find("![](data:image")
      < md2.find("题干后半段。"))
check("Markdown：标记不残留", "![[img:1]]" not in md2)

# ---------------- Markdown（旧数据，无标记） ----------------
md3 = export_markdown(view, include_images=True, answers_at_end=False)
check("Markdown（旧数据无标记）：图片仍在文字之后",
      -1 < md3.find("题干三：两组图都有") < md3.find("![](data:image"))

# ---------------- Word ----------------
from src.config import EXPORT_DIR  # noqa: E402

out = EXPORT_DIR / "_verify_export.docx"
export_word(view, out, include_images=True, answers_at_end=False)
check("Word：导出文件生成成功", out.is_file() and out.stat().st_size > 0,
      f"{out.stat().st_size if out.is_file() else 0} bytes")

from docx import Document  # noqa: E402

doc = Document(str(out))
n_pics = sum(1 for r in doc.part.rels.values() if "image" in r.reltype)
check("Word：两张图片都插入了文档", n_pics == 2, f"图片数={n_pics}")

body = "\n".join(p.text for p in doc.paragraphs)
check("Word：解析文字仍在文档里", "解析三" in body)

# 带标记的版本：图片要插在两段文字之间
out2 = EXPORT_DIR / "_verify_export_inline.docx"
export_word(view2, out2, include_images=True, answers_at_end=False)
doc2 = Document(str(out2))
paras = [p.text for p in doc2.paragraphs]
try:
    i_head = next(i for i, t in enumerate(paras) if "题干前半段。" in t)
    i_tail = next(i for i, t in enumerate(paras) if "题干后半段。" in t)
    n_pics2 = sum(1 for r in doc2.part.rels.values() if "image" in r.reltype)
    check("Word：带标记时两张图都插入", n_pics2 == 2, f"图片数={n_pics2}")

    def has_pic(p):
        return "<w:drawing" in p._p.xml or "<w:pict" in p._p.xml

    between = [i for i in range(i_head + 1, i_tail) if has_pic(doc2.paragraphs[i])]
    check("Word：图片段位于两段文字之间（不是堆在末尾）",
          bool(between), f"题干段下标 head={i_head} tail={i_tail} 图片段={between}")
except StopIteration:
    check("Word：带标记时题干两段文字都在文档里", False, f"{paras[:6]}")

# 答案放卷末的情况
html_end = export_html(view, include_images=True, answers_at_end=True)
check("HTML（答案放卷末）：图片总数仍为 2", html_end.count("data:image") == 2,
      f"count={html_end.count('data:image')}")
check("HTML（答案放卷末）：卷末答案区带上了解析图",
      html_end.rfind("参考答案与解析") < html_end.rfind("data:image"))

# ---------------- 孤儿图片清理 ----------------
from src.services import backup_service as bk  # noqa: E402

orphans = {o["rel_path"] for o in bk.find_orphan_images()}
check("被解析引用的图片不会被当成孤儿图片",
      "images/both_analysis.png" not in orphans and "images/both_stem.png" not in orphans,
      f"orphans={sorted(orphans)}")

# 「批量导入」草稿里还没确认入库的图片不算孤儿：老师可能正开着导入页在编辑，
# 扫描时把它们列出来等人点"全部删除"是很危险的。
from src.services import import_service as imp  # noqa: E402
from src.config import DATA_DIR  # noqa: E402

stage_dir = imp.staging_subdir("selftest_leftover")
stage_dir.mkdir(parents=True, exist_ok=True)
(stage_dir / "draft_only.png").write_bytes(b"x")
orphans2 = {o["rel_path"] for o in bk.find_orphan_images()}
check("导入草稿暂存目录里的图片不会被当成孤儿图片",
      not any("_import_staging" in rel for rel in orphans2),
      f"orphans={sorted(r for r in orphans2 if 'staging' in r)}")
check("草稿图片确实在磁盘上（不是靠目录不存在蒙对的）",
      (DATA_DIR / "images" / "_import_staging" / "selftest_leftover"
       / "draft_only.png").is_file())
imp.clear_staging()

print("\n================ 汇总 ================")
bad = [r for r in results if not r[1]]
for name, ok, detail in results:
    print(f"{'PASS' if ok else 'FAIL'}  {name}")
print(f"\n{len(results) - len(bad)}/{len(results)} 通过")
sys.exit(1 if bad else 0)
